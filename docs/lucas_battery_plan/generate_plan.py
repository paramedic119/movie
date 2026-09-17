# -*- coding: utf-8 -*-
"""犬山市 自動式心マッサージ器(LUCAS3) バッテリ更新計画 Excel 生成スクリプト.

元資料: 「犬山市自動式心マッサージ器 保守点検及び本体、バッテリ更新計画 (2026.5修正)」
本スクリプトは R9 以降の予備バッテリ購入数を精査した 3年維持版 / 4年維持版 を出力する。
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

JP = "ＭＳ Ｐゴシック"
YEARS = list(range(3, 20))              # R3 - R19
FIRST_COL = 3                           # C列 = R3
LAST_COL = FIRST_COL + len(YEARS) - 1   # S列 = R19
PLAN_COL = FIRST_COL + (9 - 3)          # I列 = R9 (計画開始)
TOTAL_COL = LAST_COL + 1                # T列 = 合計
VEHICLES = ["救急犬山1", "救急犬山2", "救急犬山3", "救急犬山4"]

# ---------------------------------------------------------------- 色・罫線
C_TITLE = "1F3864"
F_HEAD = PatternFill("solid", fgColor="D9E2F3")
F_YEARHEAD = PatternFill("solid", fgColor="BDD7EE")
F_JISSEKI = PatternFill("solid", fgColor="F2F2F2")   # 実績年度
F_RENEW = PatternFill("solid", fgColor="808080")     # 新規配備・機器更新
F_HOSHU = PatternFill("solid", fgColor="FFFFFF")     # 保守点検
F_BUY = PatternFill("solid", fgColor="FFE699")       # 予備購入
F_SWAP = PatternFill("solid", fgColor="E2EFDA")      # 車両間入替
F_SUM = PatternFill("solid", fgColor="FFF2CC")
F_OK = PatternFill("solid", fgColor="E2EFDA")
F_INPUT = PatternFill("solid", fgColor="FFF2CC")
F_NOTE = PatternFill("solid", fgColor="FBE5D6")

thin = Side(style="thin", color="A6A6A6")
med = Side(style="medium", color="404040")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def cell(ws, r, c, v=None, *, bold=False, size=10, fill=None, align="center",
         wrap=False, fmt=None, color=None, border=True, vert="center"):
    cl = ws.cell(row=r, column=c)
    if v is not None:
        cl.value = v
    cl.font = Font(name=JP, size=size, bold=bold, color=color)
    cl.alignment = Alignment(horizontal=align, vertical=vert, wrap_text=wrap)
    if fill:
        cl.fill = fill
    if border:
        cl.border = BOX
    if fmt:
        cl.number_format = fmt
    return cl


def title_block(ws, title, subtitle=None, width=TOTAL_COL):
    cell(ws, 1, 1, title, bold=True, size=14, align="left",
         color=C_TITLE, border=False)
    if subtitle:
        cell(ws, 2, 1, subtitle, size=9, align="left", border=False)


# ================================================================ 台帳データ
# (year, vehicle) -> (作業テキスト, 種別, 搭載バッテリ[新, 旧], 購入数)
# 種別: renew=新規配備/機器更新, hoshu=保守点検, buy=予備購入, swap=車両間入替, "" = 変化なし

def ledger_3():
    """3年維持版: 導入年度＋2年度以内のバッテリを1個以上搭載."""
    E = {}
    # ---- 実績 (R3-R8)
    E[(3, "救急犬山2")] = ("新規配備（電池2個付属）", "renew")
    E[(3, "救急犬山3")] = ("新規配備（電池2個付属）", "renew")
    E[(3, "救急犬山4")] = ("新規配備（電池2個付属）", "renew")
    E[(6, "救急犬山4")] = ("保守点検（新品1個）", "hoshu")
    E[(7, "救急犬山2")] = ("保守点検（新品1個）", "hoshu")
    E[(8, "救急犬山3")] = ("保守点検（新品1個）", "hoshu")
    # ---- 計画 (R9-R19)
    E[(9, "救急犬山1")] = ("旧救2本体を引継（電池R7・R3も引継）", "renew")
    E[(9, "救急犬山2")] = ("機器更新（車両更新・電池2個付属）\n新品1個を救4のR6と交換", "renew")
    E[(9, "救急犬山4")] = ("救2の新品1個とR6を交換", "swap")
    E[(10, "救急犬山1")] = ("予備1個購入→搭載（R3降ろす）", "buy")
    E[(11, "救急犬山3")] = ("予備1個購入→搭載（R3降ろす）", "buy")
    E[(12, "救急犬山3")] = ("機器更新（車両更新・電池2個付属）\n新品1個を救2へ／旧本体R11を救4へ", "renew")
    E[(12, "救急犬山2")] = ("救3更新の新品1個を受領（R6降ろす）", "swap")
    E[(12, "救急犬山4")] = ("旧救3本体のR11を受領（R3降ろす）", "swap")
    E[(13, "救急犬山2")] = ("保守点検（新品1個）→救1へ", "hoshu")
    E[(13, "救急犬山1")] = ("救2保守点検の新品を搭載（R7降ろす）", "swap")
    E[(14, "救急犬山4")] = ("機器更新（車両更新・電池2個付属）\n新品1個を救3へ／旧本体R11を継続搭載", "renew")
    E[(14, "救急犬山3")] = ("救4更新の新品1個を受領（R8降ろす）", "swap")
    E[(15, "救急犬山2")] = ("予備1個購入→搭載（R9降ろす）", "buy")
    E[(16, "救急犬山3")] = ("保守点検（新品1個）→救1へ", "hoshu")
    E[(16, "救急犬山1")] = ("救3保守点検の新品を搭載（R10降ろす）", "swap")
    E[(17, "救急犬山4")] = ("保守点検（新品1個）", "hoshu")
    E[(17, "救急犬山3")] = ("予備1個購入→搭載（R12降ろす）", "buy")
    E[(18, "救急犬山2")] = ("予備1個購入→搭載（R12降ろす）", "buy")
    E[(19, "救急犬山1")] = ("予備1個購入→搭載（R13降ろす）", "buy")

    B = {  # 年度末時点の搭載バッテリ [新しい方, 古い方]
        "救急犬山1": {9: (7, 3), 10: (10, 7), 13: (13, 10), 16: (16, 13), 19: (19, 16)},
        "救急犬山2": {3: (3, 3), 7: (7, 3), 9: (9, 6), 12: (12, 9), 15: (15, 12), 18: (18, 15)},
        "救急犬山3": {3: (3, 3), 8: (8, 3), 11: (11, 8), 12: (12, 8), 14: (14, 12), 17: (17, 14)},
        "救急犬山4": {3: (3, 3), 6: (6, 3), 9: (9, 3), 12: (11, 9), 14: (14, 11), 17: (17, 14)},
    }
    BUY = {(10, "救急犬山1"): 1, (11, "救急犬山3"): 1, (15, "救急犬山2"): 1,
           (17, "救急犬山3"): 1, (18, "救急犬山2"): 1, (19, "救急犬山1"): 1}
    return E, B, BUY


def ledger_4():
    """4年維持版: 導入年度＋3年度以内のバッテリを1個以上搭載."""
    E = {}
    E[(3, "救急犬山2")] = ("新規配備（電池2個付属）", "renew")
    E[(3, "救急犬山3")] = ("新規配備（電池2個付属）", "renew")
    E[(3, "救急犬山4")] = ("新規配備（電池2個付属）", "renew")
    E[(6, "救急犬山4")] = ("保守点検（新品1個）", "hoshu")
    E[(7, "救急犬山2")] = ("保守点検（新品1個）", "hoshu")
    E[(8, "救急犬山3")] = ("保守点検（新品1個）", "hoshu")
    E[(9, "救急犬山1")] = ("旧救2本体を引継（電池R7・R3も引継）", "renew")
    E[(9, "救急犬山2")] = ("機器更新（車両更新・電池2個付属）", "renew")
    E[(10, "救急犬山2")] = ("新品1個を救4のR6と交換", "swap")
    E[(10, "救急犬山4")] = ("救2の新品1個とR6を交換", "swap")
    E[(11, "救急犬山1")] = ("予備1個購入→搭載（R3降ろす）", "buy")
    E[(12, "救急犬山3")] = ("機器更新（車両更新・電池2個付属）\n旧本体R8を救4へ", "renew")
    E[(12, "救急犬山4")] = ("旧救3本体のR8を受領（R3降ろす）", "swap")
    E[(13, "救急犬山2")] = ("保守点検（新品1個）（R6降ろす）", "hoshu")
    E[(13, "救急犬山3")] = ("新品R12と救4のR9を交換", "swap")
    E[(13, "救急犬山4")] = ("救3のR12とR9を交換", "swap")
    E[(14, "救急犬山4")] = ("機器更新（車両更新・電池2個付属）\n新品1個を救1へ／旧本体R12を継続搭載", "renew")
    E[(14, "救急犬山1")] = ("救4更新の新品1個を受領（R7降ろす）", "swap")
    E[(16, "救急犬山3")] = ("保守点検（新品1個）（R9降ろす）", "hoshu")
    E[(17, "救急犬山4")] = ("保守点検（新品1個）→救2へ", "hoshu")
    E[(17, "救急犬山2")] = ("救4保守点検の新品を搭載（R9降ろす）", "swap")
    E[(18, "救急犬山1")] = ("予備1個購入→搭載（R11降ろす）", "buy")
    E[(18, "救急犬山4")] = ("予備1個購入→搭載（R12降ろす）", "buy")

    B = {
        "救急犬山1": {9: (7, 3), 11: (11, 7), 14: (14, 11), 18: (18, 14)},
        "救急犬山2": {3: (3, 3), 7: (7, 3), 9: (9, 9), 10: (9, 6), 13: (13, 9), 17: (17, 13)},
        "救急犬山3": {3: (3, 3), 8: (8, 3), 12: (12, 12), 13: (12, 9), 16: (16, 12)},
        "救急犬山4": {3: (3, 3), 6: (6, 3), 10: (9, 3), 12: (9, 8), 13: (12, 8), 14: (14, 12), 18: (18, 14)},
    }
    BUY = {(11, "救急犬山1"): 1, (18, "救急犬山1"): 1, (18, "救急犬山4"): 1}
    return E, B, BUY


def expand(bat):
    """飛び飛びの台帳を全年度に展開する."""
    out, cur = {}, None
    for y in YEARS:
        if y in bat:
            cur = bat[y]
        out[y] = cur
    return out


# ================================================================ 年表シート
def build_plan_sheet(wb, name, keep_years, ledger):
    E, B, BUY = ledger
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    title_block(
        ws,
        f"犬山市　自動式心マッサージ器（LUCAS3）　バッテリ更新計画　【{keep_years}年維持版】",
        "元資料「保守点検及び本体、バッテリ更新計画（2026.5修正）」の機器更新・保守点検スケジュールはそのまま、"
        "R9以降の予備バッテリ購入数のみ精査したもの。",
    )
    cell(ws, 3, 1, "判定基準（年）", bold=True, size=9, fill=F_HEAD, align="center")
    cell(ws, 3, 2, keep_years, bold=True, size=11, fill=F_INPUT)
    cell(ws, 3, 3, f"…各車とも「導入年度＋{keep_years - 1}年度以内」に導入したバッテリを1個以上搭載していればOK",
         size=9, align="left", border=False)

    hr = 5                         # 見出し行
    cell(ws, hr, 1, "車両", bold=True, fill=F_YEARHEAD)
    cell(ws, hr, 2, "項目", bold=True, fill=F_YEARHEAD)
    for i, y in enumerate(YEARS):
        c = FIRST_COL + i
        cell(ws, hr, c, y, bold=True, fill=F_YEARHEAD, fmt='"R"0')
        cell(ws, hr + 1, c, "実績" if y < 9 else "計画", size=8,
             fill=F_JISSEKI if y < 9 else F_HEAD)
    cell(ws, hr, TOTAL_COL, "R9〜R19\n合計", bold=True, fill=F_YEARHEAD, wrap=True)
    cell(ws, hr, 1, "車両", bold=True, fill=F_YEARHEAD)
    cell(ws, hr + 1, 1, "区分", bold=True, size=8, fill=F_YEARHEAD)
    cell(ws, hr + 1, 2, "", fill=F_YEARHEAD)
    cell(ws, hr + 1, TOTAL_COL, "", fill=F_YEARHEAD)

    fills = {"renew": F_RENEW, "hoshu": F_HOSHU, "buy": F_BUY, "swap": F_SWAP}
    rows = {}
    r = hr + 2
    for v in VEHICLES:
        bat = expand(B[v])
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 3, end_column=1)
        cell(ws, r, 1, f"{v}\n(LUCAS3)", bold=True, fill=F_HEAD, wrap=True)
        for rr in range(r, r + 4):
            cell(ws, rr, 1, fill=F_HEAD)
        cell(ws, r, 2, "機器イベント・作業", size=8, fill=F_HEAD, align="left", wrap=True)
        cell(ws, r + 1, 2, "搭載バッテリ（新）", size=8, fill=F_HEAD, align="left")
        cell(ws, r + 2, 2, "搭載バッテリ（旧）", size=8, fill=F_HEAD, align="left")
        cell(ws, r + 3, 2, "予備バッテリ購入数", size=8, fill=F_HEAD, align="left")
        rows[v] = r
        for i, y in enumerate(YEARS):
            c = FIRST_COL + i
            base = F_JISSEKI if y < 9 else None
            ev = E.get((y, v))
            if ev:
                txt, kind = ev
                fc = "FFFFFF" if kind == "renew" else None
                cell(ws, r, c, txt, size=7, fill=fills[kind], wrap=True,
                     bold=(kind == "renew"), color=fc)
            else:
                cell(ws, r, c, fill=base)
            b = bat[y]
            cell(ws, r + 1, c, b[0] if b else None, size=9, fmt='"R"0',
                 fill=base, bold=True)
            cell(ws, r + 2, c, b[1] if b else None, size=9, fmt='"R"0', fill=base)
            cell(ws, r + 3, c, BUY.get((y, v), 0), size=9, fill=F_BUY if BUY.get((y, v)) else base)
        for k, lab in enumerate(["", "", "", ""]):
            cell(ws, r + k, TOTAL_COL, fill=F_SUM if k == 3 else None)
        L, R = get_column_letter(PLAN_COL), get_column_letter(LAST_COL)
        cell(ws, r + 3, TOTAL_COL, f"=SUM({L}{r + 3}:{R}{r + 3})", bold=True, fill=F_SUM)
        r += 4

    # ---- 集計・判定
    sum_r = r
    cell(ws, sum_r, 1, "予備バッテリ", bold=True, fill=F_SUM)
    cell(ws, sum_r, 2, "購入数（年度計）", bold=True, size=9, fill=F_SUM, align="left")
    cell(ws, sum_r + 1, 1, "", fill=F_SUM)
    cell(ws, sum_r + 1, 2, "購入数（累計）", size=9, fill=F_SUM, align="left")
    cell(ws, sum_r + 2, 1, "判定", bold=True, fill=F_OK)
    cell(ws, sum_r + 2, 2, "全車に耐用年数内の電池1個以上", size=8, fill=F_OK, align="left", wrap=True)

    for i, y in enumerate(YEARS):
        c = get_column_letter(FIRST_COL + i)
        parts = "+".join(f"{c}{rows[v] + 3}" for v in VEHICLES)
        cell(ws, sum_r, FIRST_COL + i, f"={parts}", bold=True, size=10, fill=F_SUM)
        if i == 0:
            cell(ws, sum_r + 1, FIRST_COL + i, f"={c}{sum_r}", size=9, fill=F_SUM)
        else:
            p = get_column_letter(FIRST_COL + i - 1)
            cell(ws, sum_r + 1, FIRST_COL + i, f"={p}{sum_r + 1}+{c}{sum_r}", size=9, fill=F_SUM)
        conds = ",".join(
            f"MAX({c}{rows[v] + 1}:{c}{rows[v] + 2})>={c}${hr}-$B$3+1" for v in VEHICLES)
        cell(ws, sum_r + 2, FIRST_COL + i,
             f'=IF({c}${hr}<9,"－",IF(AND({conds}),"○","×"))', bold=True, fill=F_OK)
    L, R = get_column_letter(PLAN_COL), get_column_letter(LAST_COL)
    cell(ws, sum_r, TOTAL_COL, f"=SUM({L}{sum_r}:{R}{sum_r})", bold=True, size=11, fill=F_SUM)
    cell(ws, sum_r + 1, TOTAL_COL, fill=F_SUM)
    cell(ws, sum_r + 2, TOTAL_COL,
         f'=IF(COUNTIF({L}{sum_r + 2}:{R}{sum_r + 2},"×")=0,"全年度OK","要確認")',
         bold=True, fill=F_OK)

    # ---- 凡例
    lr = sum_r + 4
    cell(ws, lr, 1, "凡例", bold=True, size=9, border=False, align="left")
    legend = [("新規配備／機器更新（新品電池2個付属）", F_RENEW, "FFFFFF"),
              ("保守点検（新品電池1個付属）", F_HOSHU, None),
              ("予備バッテリ購入", F_BUY, None),
              ("車両間でのバッテリ入替（購入なし）", F_SWAP, None),
              ("実績年度（R3〜R8）", F_JISSEKI, None)]
    for k, (t, f, col) in enumerate(legend):
        cell(ws, lr + 1 + k, 1, "", fill=f)
        cell(ws, lr + 1 + k, 2, t, size=9, align="left", border=False)
    nr = lr + 1 + len(legend) + 1
    notes = [
        "※「搭載バッテリ（新）（旧）」は、その年度末時点で各車に積んでいる2個の電池の導入年度。",
        "※ 購入数・判定はすべて数式。判定基準（B3）を書き換えると全年度の判定が再計算される（搭載計画そのものは自動では変わらない）。",
        f"※ 判定は R9 以降のみ（R3〜R8 は実績のため「－」）。B3={keep_years} のとき全年度「○」。",
        "※ R9の予備バッテリ購入は0個。救2の機器更新で新品2個が入るため。",
    ]
    for k, t in enumerate(notes):
        cell(ws, nr + k, 1, t, size=9, align="left", border=False)

    # ---- 体裁
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 20
    for i in range(len(YEARS)):
        ws.column_dimensions[get_column_letter(FIRST_COL + i)].width = 14
    ws.column_dimensions[get_column_letter(TOTAL_COL)].width = 10
    for v in VEHICLES:
        ws.row_dimensions[rows[v]].height = 52
    ws.row_dimensions[hr].height = 20
    ws.row_dimensions[sum_r + 2].height = 22
    ws.freeze_panes = ws.cell(row=hr + 2, column=FIRST_COL)
    return ws


# ================================================================ 比較まとめ
SCENARIOS = [
    ("元計画（2026.5修正）", {9: 3, 12: 3, 15: 3, 18: 3}, "3年ごとに一律3個を計上"),
    ("3年維持版（車両間入替あり）", {10: 1, 11: 1, 15: 1, 17: 1, 18: 1, 19: 1},
     "推奨案。常に「導入年度＋2年度以内」の電池を1個以上確保"),
    ("4年維持版（車両間入替あり）", {11: 1, 18: 2},
     "推奨案。常に「導入年度＋3年度以内」の電池を1個以上確保"),
    ("参考：3年維持版（入替なし）", {9: 1, 10: 1, 11: 1, 12: 2, 13: 1, 15: 1, 16: 2, 19: 3},
     "各車に積みっぱなしで運用した場合"),
    ("参考：4年維持版（入替なし）", {10: 1, 11: 1, 15: 1, 17: 1, 19: 1},
     "各車に積みっぱなしで運用した場合"),
]


def build_summary(wb):
    ws = wb.create_sheet("比較まとめ", 0)
    ws.sheet_view.showGridLines = False
    title_block(
        ws,
        "犬山市　自動式心マッサージ器（LUCAS3）　予備バッテリ購入数の比較",
        "元資料「保守点検及び本体、バッテリ更新計画（2026.5修正）」の R9「予備3個購入」を見直したもの。"
        "機器更新・保守点検のスケジュールは元資料のまま変更していない。",
    )

    hr = 4
    cell(ws, hr, 1, "区分", bold=True, fill=F_YEARHEAD)
    for i, y in enumerate(YEARS):
        cell(ws, hr, FIRST_COL + i, y, bold=True, fmt='"R"0',
             fill=F_JISSEKI if y < 9 else F_YEARHEAD)
    cell(ws, hr, TOTAL_COL, "R9〜R19\n合計", bold=True, fill=F_YEARHEAD, wrap=True)
    cell(ws, hr, TOTAL_COL + 1, "考え方", bold=True, fill=F_YEARHEAD)
    ws.merge_cells(start_row=hr - 1, start_column=FIRST_COL, end_row=hr - 1, end_column=PLAN_COL - 1)
    cell(ws, hr - 1, FIRST_COL, "実績（購入なし）", size=8, fill=F_JISSEKI)
    ws.merge_cells(start_row=hr - 1, start_column=PLAN_COL, end_row=hr - 1, end_column=LAST_COL)
    cell(ws, hr - 1, PLAN_COL, "計画期間", size=8, fill=F_HEAD)

    r = hr + 1
    rowmap = {}
    for label, data, memo in SCENARIOS:
        ref = label.startswith("参考")
        f = F_SUM if not ref and "版" in label else None
        cell(ws, r, 1, label, bold=not ref, size=10, align="left",
             fill=F_HEAD if not ref else F_JISSEKI)
        for i, y in enumerate(YEARS):
            v = data.get(y, 0)
            cell(ws, r, FIRST_COL + i, v, size=10, bold=bool(v) and not ref,
                 fill=(F_BUY if v and not ref else (F_JISSEKI if y < 9 else None)))
        L, R = get_column_letter(PLAN_COL), get_column_letter(LAST_COL)
        cell(ws, r, TOTAL_COL, f"=SUM({L}{r}:{R}{r})", bold=True, size=11, fill=F_SUM)
        cell(ws, r, TOTAL_COL + 1, memo, size=9, align="left", wrap=True)
        rowmap[label] = r
        r += 1

    # ---- 差引
    base = rowmap["元計画（2026.5修正）"]
    T = get_column_letter(TOTAL_COL)
    for label, key in [("3年維持版（車両間入替あり）", "3年維持版"), ("4年維持版（車両間入替あり）", "4年維持版")]:
        rr = rowmap[label]
        cell(ws, r, 1, f"元計画との差引（{key}）", bold=True, size=10, align="left", fill=F_NOTE)
        for i, y in enumerate(YEARS):
            c = get_column_letter(FIRST_COL + i)
            cell(ws, r, FIRST_COL + i, f"={c}{rr}-{c}{base}", size=10, fill=F_NOTE)
        cell(ws, r, TOTAL_COL, f"={T}{rr}-{T}{base}", bold=True, size=11, fill=F_NOTE)
        cell(ws, r, TOTAL_COL + 1, "マイナスが削減個数", size=9, align="left")
        r += 1

    # ---- 事業費試算
    r += 1
    cell(ws, r, 1, "バッテリ単価（税込・円）", bold=True, size=10, align="left", fill=F_HEAD)
    price = cell(ws, r, 2, 0, bold=True, size=11, fill=F_INPUT, fmt="#,##0")
    cell(ws, r, 3, "← 単価を入力すると下の事業費が自動計算されます", size=9,
         align="left", border=False)
    price_ref = "$B$" + str(r)
    r += 1
    cost_start = r
    for label in ["元計画（2026.5修正）", "3年維持版（車両間入替あり）", "4年維持版（車両間入替あり）"]:
        rr = rowmap[label]
        cell(ws, r, 1, f"事業費　{label}", size=10, align="left", fill=F_HEAD)
        for i, y in enumerate(YEARS):
            c = get_column_letter(FIRST_COL + i)
            cell(ws, r, FIRST_COL + i, f"={c}{rr}*{price_ref}", size=9, fmt="#,##0")
        cell(ws, r, TOTAL_COL, f"={T}{rr}*{price_ref}", bold=True, size=10,
             fmt="#,##0", fill=F_SUM)
        r += 1

    # ---- 結論
    r += 1
    for t, b in [
        ("【結論】", True),
        ("・R9（令和9年度）の予備バッテリ購入は、3年維持版・4年維持版のいずれも 0個。"
         "救急犬山2の機器更新（車両更新）で新品バッテリ2個が付属するため。", False),
        ("・R9〜R19の購入数合計は、元計画12個 → 3年維持版6個／4年維持版3個。", False),
        ("・上記個数は、機器更新や保守点検で入った新品・使用年数の浅い電池を車両間で入れ替えることが条件。"
         "入れ替えを行わない場合は 3年維持版12個／4年維持版5個（参考行）。", False),
        ("・本表は元資料の機器更新・保守点検スケジュールを一切変更していない。"
         "変更点は予備バッテリの購入年度と個数のみ。", False),
    ]:
        cell(ws, r, 1, t, bold=b, size=10, align="left", border=False)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=TOTAL_COL)
        r += 1

    ws.column_dimensions["A"].width = 30
    for i in range(len(YEARS)):
        ws.column_dimensions[get_column_letter(FIRST_COL + i)].width = 8
    ws.column_dimensions[get_column_letter(TOTAL_COL)].width = 11
    ws.column_dimensions[get_column_letter(TOTAL_COL + 1)].width = 46
    ws.row_dimensions[hr].height = 22
    ws.freeze_panes = ws.cell(row=hr + 1, column=FIRST_COL)
    return ws


# ================================================================ 作業手順
STEPS_3 = [
    (9, "救2の新品1個と救4のR6を交換", 0),
    (10, "1個購入 → 救1に搭載", 1),
    (11, "1個購入 → 救3に搭載", 1),
    (12, "救3更新の新品1個 → 救2へ／旧救3本体のR11 → 救4へ", 0),
    (13, "救2保守点検の新品 → 救1へ", 0),
    (14, "救4更新の新品1個 → 救3へ", 0),
    (15, "1個購入 → 救2に搭載", 1),
    (16, "救3保守点検の新品 → 救1へ", 0),
    (17, "1個購入 → 救3に搭載", 1),
    (18, "1個購入 → 救2に搭載", 1),
    (19, "1個購入 → 救1に搭載", 1),
]
STEPS_4 = [
    (9, "作業なし（救2の機器更新で新品2個が入るため）", 0),
    (10, "救2の新品1個と救4のR6を交換", 0),
    (11, "1個購入 → 救1に搭載", 1),
    (12, "旧救3本体のR8 → 救4へ", 0),
    (13, "救3のR12と救4のR9を交換", 0),
    (14, "救4更新の新品1個 → 救1へ", 0),
    (17, "救4保守点検の新品 → 救2へ", 0),
    (18, "2個購入 → 救1・救4に搭載", 2),
]


def build_steps(wb):
    ws = wb.create_sheet("作業手順")
    ws.sheet_view.showGridLines = False
    title_block(ws, "年度ごとのバッテリ入替・購入作業",
                "各年度に行う作業。購入がない年度も入替作業が必要なので注意。")
    r = 4
    for name, steps, keep in [("3年維持版", STEPS_3, 3), ("4年維持版", STEPS_4, 4)]:
        cell(ws, r, 1, f"【{name}】（導入年度＋{keep - 1}年度以内の電池を1個以上維持）",
             bold=True, size=11, align="left", fill=F_HEAD)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
        for h, w in [("年度", 1), ("実施内容", 2), ("購入数", 3), ("備考", 4)]:
            cell(ws, r, w, h, bold=True, fill=F_YEARHEAD)
        r += 1
        top = r
        for y, txt, n in steps:
            cell(ws, r, 1, y, bold=True, fmt='"R"0',
                 fill=F_BUY if n else None)
            cell(ws, r, 2, txt, size=10, align="left", wrap=True)
            cell(ws, r, 3, n, bold=bool(n), fill=F_BUY if n else None)
            memo = "購入あり" if n else ("購入なし" if "作業なし" in txt else "入替のみ（購入なし）")
            cell(ws, r, 4, memo, size=9)
            r += 1
        cell(ws, r, 1, "合計", bold=True, fill=F_SUM)
        cell(ws, r, 2, "", fill=F_SUM)
        cell(ws, r, 3, f"=SUM(C{top}:C{r - 1})", bold=True, size=11, fill=F_SUM)
        cell(ws, r, 4, "", fill=F_SUM)
        r += 3
    notes = [
        "※ R9は両案とも購入0個。救急犬山2の機器更新で新品バッテリ2個が入るため。",
        "※ 「交換」は2台の車の間で電池を入れ替えること。「→○○へ」は付属してきた新品を別の車に回すこと。",
        "※ 機器更新で外した旧本体に付いていた電池は、新本体で使用できるか事前にメーカーへ要確認。",
    ]
    for k, t in enumerate(notes):
        cell(ws, r + k, 1, t, size=9, align="left", border=False)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 22
    return ws


# ================================================================ 前提条件
def build_assumptions(wb):
    ws = wb.create_sheet("前提条件")
    ws.sheet_view.showGridLines = False
    title_block(ws, "前提条件・元資料の内容",
                "本計画の計算根拠。数値を見直す場合はここを確認すること。")
    r = 4

    def section(t):
        nonlocal r
        cell(ws, r, 1, t, bold=True, size=11, align="left", fill=F_HEAD)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        r += 1

    def line(a, b="", c=""):
        nonlocal r
        cell(ws, r, 1, a, size=10, align="left", wrap=True)
        cell(ws, r, 2, b, size=10, align="left", wrap=True)
        cell(ws, r, 3, c, size=10, align="left", wrap=True)
        r += 1

    section("1．機器及び付属品の耐用期間（元資料より）")
    line("LUCAS3本体", "8年")
    line("バッテリ", "3年〜4年", "使用頻度と実際は3年目にバッテリエラー表示")
    r += 1

    section("2．バッテリの供給元（付属する新品の個数）")
    line("新規配備", "新品2個", "")
    line("機器更新（車両更新）", "新品2個", "新規配備と同じ")
    line("保守点検", "新品1個", "保守点検プランは「基本保守点検」")
    line("予備バッテリ購入", "購入した個数", "不足分のみ購入")
    r += 1

    section("3．判定ルール")
    line("各車の搭載数", "2個", "")
    line("3年維持版", "2個のうち1個が「導入年度＋2年度以内」ならOK", "")
    line("4年維持版", "2個のうち1個が「導入年度＋3年度以内」ならOK", "")
    line("購入数の考え方", "配備台数 − 直近N年度に入った付属の新品数 − 同じ期間の購入数（0未満は0）",
         "期限が来た年度に、足りない分だけ購入")
    line("延命の扱い", "交換推奨は3〜4年だが、新品が来るまでは使用を継続（延命）", "")
    r += 1

    section("4．元資料（2026.5修正）の機器スケジュール　※本計画では変更していない")
    cell(ws, r, 1, "車両", bold=True, fill=F_YEARHEAD)
    cell(ws, r, 2, "新規配備", bold=True, fill=F_YEARHEAD)
    cell(ws, r, 3, "保守点検", bold=True, fill=F_YEARHEAD)
    cell(ws, r, 4, "機器更新（車両更新）", bold=True, fill=F_YEARHEAD)
    r += 1
    for v, a, b, c in [
        ("救急犬山1", "―（R9に旧救2本体を引継）", "―", "―（元資料に記載なし）"),
        ("救急犬山2", "R3", "R7、R13", "R9"),
        ("救急犬山3", "R3", "R8、R16", "R12"),
        ("救急犬山4", "R3", "R6、R17", "R14"),
    ]:
        cell(ws, r, 1, v, bold=True, size=10, fill=F_HEAD)
        cell(ws, r, 2, a, size=10)
        cell(ws, r, 3, b, size=10)
        cell(ws, r, 4, c, size=10)
        r += 1
    r += 1

    section("5．元資料の備考")
    for t in ["LUCAS本体の更新は、救急車車両更新(11年)併せて実施する。",
              "保守点検は出動件数の多い車両の機器から順に実施する。",
              "保守点検プランはLUCAS3メンテナンスサービス表の「基本保守点検」とする。",
              "バッテリは保守点検時に新品1個が付属するが、保守点検とは別に予備バッテリの購入を行い、"
              "常に耐用年数内のバッテリが1つは積載されている状態で運用を行う。"]:
        cell(ws, r, 1, "・" + t, size=10, align="left", wrap=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 16
        r += 1
    r += 1

    section("6．本計画で置いた前提")
    for t in ["R9に、旧救急犬山2本体（R7とR3のバッテリ付き）を救急犬山1へ積載する。",
              "機器更新で外した旧本体のバッテリは、他の車で再利用する。",
              "救急犬山1の本体更新は元資料に記載がないため、本計画には含めていない。",
              "R3〜R8は実績として扱い、購入計画はR9（令和9年度）以降を対象とする。"]:
        cell(ws, r, 1, "・" + t, size=10, align="left", wrap=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        r += 1
    r += 1

    section("7．要確認事項")
    for t in ["旧本体から外したバッテリが新本体でそのまま使用できるか（メーカー確認）。",
              "救急犬山1の本体は R3配備（旧救2）のため、耐用期間8年は R11 に到達する。"
              "R11以降の本体の取扱い（救3・救4の旧本体を回すか等）は元資料に記載がなく、別途検討が必要。",
              "車両間でバッテリを入れ替える運用が可能か（管理簿・資器材管理上の取扱い）。",
              "入替を行わない場合は、比較まとめシートの「参考」行（3年版12個／4年版5個）となる。"]:
        cell(ws, r, 1, "・" + t, size=10, align="left", wrap=True, fill=F_NOTE)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        ws.row_dimensions[r].height = 28
        r += 1

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 26
    return ws


# ================================================================ main
def main():
    wb = Workbook()
    wb.remove(wb.active)
    build_plan_sheet(wb, "3年維持版", 3, ledger_3())
    build_plan_sheet(wb, "4年維持版", 4, ledger_4())
    build_steps(wb)
    build_assumptions(wb)
    build_summary(wb)                      # index 0 に挿入
    wb.move_sheet("比較まとめ", offset=0)
    wb.active = 0
    out = "犬山市_LUCAS3バッテリ更新計画_R9修正.xlsx"
    wb.save(out)
    print("saved:", out)
    print("sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
